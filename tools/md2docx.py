# -*- coding: utf-8 -*-
"""
Markdown → Word (.docx) 转换脚本
用法: python tools/md2docx.py <输入.md> <输出.docx>

支持语法: #/##/### 标题、``` 代码块、表格(|...|)、- [ ] 复选框、
- 列表、> 引用、**加粗**、`行内代码`。
中文文档排版: 正文宋体、标题微软雅黑、代码 Consolas。
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"
CODE_FONT = "Consolas"
CODE_BG = "F2F2F2"


def set_run_font(run, name, size_pt, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), BODY_FONT if name == BODY_FONT else name)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_inline_runs(paragraph, text, base_font, size, bold_all=False):
    """解析 **加粗** 与 `代码` 行内标记并写入 runs"""
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            set_run_font(run, base_font, size, bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            set_run_font(run, CODE_FONT, size - 0.5)
        else:
            run = paragraph.add_run(tok)
            set_run_font(run, base_font, size, bold=bold_all)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Pt(12)
        run = p.add_run(line if line else " ")
        set_run_font(run, CODE_FONT, 9)
        shade_paragraph(p, CODE_BG)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            add_inline_runs(p, cell, BODY_FONT, 10, bold_all=(i == 0))
    doc.add_paragraph()


def main():
    if len(sys.argv) != 3:
        print("用法: python tools/md2docx.py <输入.md> <输出.docx>")
        sys.exit(1)
    src, dst = sys.argv[1], sys.argv[2]

    with open(src, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)

        elif line.startswith("### "):
            p = doc.add_heading(level=3)
            add_inline_runs(p, line[4:], HEADING_FONT, 13, bold_all=True)

        elif line.startswith("## "):
            p = doc.add_heading(level=2)
            add_inline_runs(p, line[3:], HEADING_FONT, 15, bold_all=True)

        elif line.startswith("# "):
            p = doc.add_heading(level=1)
            add_inline_runs(p, line[2:], HEADING_FONT, 17, bold_all=True)

        elif line.startswith("|"):
            rows = []
            while i < n and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        elif line.startswith("- [ ] "):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Pt(12)
            pf.space_after = Pt(2)
            run = p.add_run("☐ ")
            set_run_font(run, BODY_FONT, 11, bold=True)
            add_inline_runs(p, line[6:], BODY_FONT, 11, bold_all=True)

        elif line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("• ")
            set_run_font(run, BODY_FONT, 11)
            add_inline_runs(p, line[2:], BODY_FONT, 11)

        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            add_inline_runs(p, line[2:], BODY_FONT, 10)
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

        elif line.strip() == "---":
            pass

        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, line, BODY_FONT, 11)

        i += 1

    doc.save(dst)
    print(f"已生成: {dst}")


if __name__ == "__main__":
    main()
