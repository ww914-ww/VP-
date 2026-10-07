# -*- coding: utf-8 -*-
"""Generate VP运维工具 灰度更新功能 项目规划方案书 V2.0 as Word document.

V2.0 修订（2026-09-29）：
- 新增 1.1 旧更新故障复盘（"下载无效+反复提示"三条根因）与 Launcher 残留风险
- 新增 阶段0：更新链路可靠性修复（版本真相源 / manifest 校验 / 废弃 bat / 失败计数）
- 新增 可选灰度：本地 GrayOptIn 开关 + 云端 requireOptIn 策略
- 检查时机定稿：Launcher 启动时检查（6s 超时）+ 主程序运行中定时后台检查
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# -- page setup --
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# -- styles --
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
font.size = Pt(10.5)

for level in range(1, 4):
    h = doc.styles['Heading %d' % level]
    h.font.name = '微软雅黑'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    h.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

FONT = '微软雅黑'
CODE_FONT = 'Consolas'
BLUE = RGBColor(0x00, 0x78, 0xd4)
GRAY = RGBColor(0x88, 0x88, 0x88)
DARK = RGBColor(0x33, 0x33, 0x33)
RED = RGBColor(0xC0, 0x39, 0x2B)


def set_font(run, name=FONT, size=None, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def p(text, bold=False, size=None, indent=0, align=None, gap=6, color=None):
    par = doc.add_paragraph()
    if align is not None:
        par.alignment = align
    par.paragraph_format.space_after = Pt(gap)
    par.paragraph_format.left_indent = Cm(indent)
    r = par.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return par


def code(text):
    for line in text.strip().split('\n'):
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.left_indent = Cm(0.5)
        r = cp.add_run(line)
        set_font(r, name=CODE_FONT, size=9, color=DARK)
    doc.add_paragraph()


def bullet(text, indent=0.5):
    bp = doc.add_paragraph(style='List Bullet')
    bp.clear()
    r = bp.add_run(text)
    set_font(r, size=10)
    bp.paragraph_format.left_indent = Cm(indent)


def tbl(headers, rows):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=9, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ''
            r = cell.paragraphs[0].add_run(val)
            set_font(r, size=9)
    doc.add_paragraph()


def heading(text, level=1):
    doc.add_heading(text, level=level)


# ===================================================================
# COVER
# ===================================================================
for _ in range(6):
    doc.add_paragraph()
p('VP运维工具', bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER, gap=12)
p('灰度更新功能 项目规划方案书', bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, gap=30)
p('方案代号：Gray Update V2 / 形态1+（SMB 更新下载 + SFTP 心跳上报 + 可选灰度）', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, gap=6, color=GRAY)
p('文档版本：V2.0（定稿）', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, gap=6, color=GRAY)
p('编制日期：2026 年 9 月 29 日', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, gap=6, color=GRAY)
doc.add_paragraph()
p('V2.0 修订说明：', bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, gap=2, color=GRAY)
p('① 新增 1.1 旧自动更新故障复盘与根因定位；② 新增"阶段0：更新链路可靠性修复"（先于灰度实施）；',
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, gap=2, color=GRAY)
p('③ 新增"可选灰度"（本地 GrayOptIn 开关 + 云端 requireOptIn 策略）；④ 检查时机定稿（启动时 + 运行中定时）。',
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, gap=2, color=GRAY)
doc.add_page_break()

# ===================================================================
# 目录
# ===================================================================
heading('目录', 1)
toc = [
    '第一章 项目概述',
    '    1.1 项目背景与现状问题（含旧更新故障复盘）',
    '    1.2 项目目标',
    '    1.3 范围与技术约束',
    '    1.4 关键决策记录',
    '第二章 初步思路与计划',
    '    2.1 总体架构（形态1+）',
    '    2.2 核心设计要点',
    '    2.3 五阶段实施计划总览',
    '第三章 实施内容与过程',
    '    3.1 阶段0：更新链路可靠性修复（先行）',
    '    3.2 阶段1：灰度判定核心库',
    '    3.3 阶段2：Launcher 集成',
    '    3.4 阶段3：主程序心跳上报（SFTP）',
    '    3.5 阶段4：运维 SOP 与闭环验证',
    '    3.6 风险识别与应对',
    '    3.7 里程碑与交付节点',
    '第四章 最终期望结果',
    '    4.1 功能验收标准',
    '    4.2 交付物清单',
    '    4.3 典型使用场景（发布演练剧本）',
]
for item in toc:
    p(item, size=11, gap=2)
doc.add_page_break()

# ===================================================================
# 第一章 项目概述
# ===================================================================
heading('第一章 项目概述', 1)

heading('1.1 项目背景与现状问题', 2)

heading('1.1.1 旧自动更新故障复盘（"下载无效 + 反复提示"根因定位）', 3)
p('VP运维工具此前在主程序内置"版本更新"Tab（v1 应用内更新器，已于 v1.0.4 删除，git a132baa），现场反馈：弹出检测到新版本的提示后，点击更新、从 SMB 云端下载了新版本，但重启后仍是旧版本，且每次打开都重复提示更新，最终弃用。经代码复盘，根因有三：', )
tbl(['#', '根因', '机理与后果'], [
    ['1', '版本真相源错误（最致命）',
     '旧更新器以 Assembly.GetExecutingAssembly().GetName().Version（程序集版本号）作为本地版本，而 AssemblyInfo.cs 写死 1.0.3.0、发新版时从不修改。即使更新安装成功，重启后程序集版本仍旧 < 云端 latest → 每次启动都判定"有新版本"→ 无限重复提示，用户观感"还是旧版本"。'],
    ['2', '安装环节是无人值守的 bat，失败无声无息',
     'update.bat 以 UTF-8（带 BOM）写入却首行 chcp 65001，老 cmd 解析出错、中文路径乱码；taskkill 后文件锁未释放导致 robocopy 移动失败；bat 中途死亡后没有任何进程把主程序拉起，也没有日志与状态记录 → 点了更新，重启后一切照旧。'],
    ['3', 'SMB 挂载不检查错误码',
     'net use Y: 未检查退出码。典型错误 1219：同一服务器（172.16.22.138）已被搬运业务以另一身份连接，Windows 同会话只允许一套身份。挂载失败后可能读到错误共享的 version.json 或下载到错误内容，表现为"下载成功但安装失败"。'],
])
p('结论：旧体系的失败不是"灰度策略"问题，而是更新链路本身不可靠。任何灰度方案若不先修复更新链路，只会把同样的故障放大到更多机台。', bold=True, color=RED)

heading('1.1.2 当前 Launcher 链路残留风险（灰度实施前必须处理）', 3)
p('现行更新体系为 Launcher.exe 前置检查（v2 方案）：启动读 version.json → SMB 复制到 versions\\.temp → update.bat 安装到 versions\\{版本}\\ → 启动主程序，带 update.status 状态机与 backup\\ 崩溃恢复。代码走查发现以下残留风险：')
tbl(['#', '风险点', '说明'], [
    ['1', '启动时实际不检查更新', 'Window_Loaded 步骤3：只要本地存在可用版本（VerifyCurrentVersion 为真）即直接启动、跳过 SMB 检查。现网机台实际上从不检查更新，灰度放量无从谈起。'],
    ['2', 'update.bat 编码地雷', 'Launcher 生成的 bat 以 GBK（Encoding.Default）写入却首行 chcp 65001，chcp 之后 cmd 按 UTF-8 解析剩余 GBK 字节，部署路径含中文时必炸——与旧 v1 故障同型。'],
    ['3', '失败后状态机死循环', 'bat 执行失败时 update.status 停留在 ready/installing，下次启动反复重装/恢复，无失败计数与人工兜底出口。'],
    ['4', '版本目录判断脆弱', 'GetLocalLatestVersion() 扫描 versions\\ 目录取最大版本号；半成品目录（有目录无 exe）会使 VerifyCurrentVersion 与 StartMainApp 判断错乱。'],
    ['5', 'Launcher 自身不可更新', '该机制只能替换 versions\\ 下的主程序文件，Launcher.exe 永远是最初人工部署的版本，灰度新特性（判定逻辑）无法触达老机台。'],
])

heading('1.1.3 全量发布风险（灰度需求原始动机）', 3)
p('latest 一经修改，所有机台在下一次启动时同时升级，属于"全量发布"。一旦新版本存在缺陷，将瞬间波及全部生产机台，且无法控制放量节奏、无法观察小范围运行效果后再推广。因此需要灰度更新能力：先小范围投放新版本，确认稳定后再逐步放量直至全量，并全程可观测、可回滚；同时现场希望灰度参与是"可选"的——愿意试新的机台可主动加入，关键机台不被强制。')

heading('1.2 项目目标', 2)
bullet('可靠更新（新，阶段0）：更新链路端到端可验证——下载有完整性校验、安装有日志与结果回报、失败有计数与自动回滚、同一版本不重复提示；彻底消除 1.1.1 的三类故障。')
bullet('灰度可控：支持白名单优先、黑名单排除、百分比放量三级灰度策略，放量进度由运维人员通过修改服务器端配置文件即可调节。')
bullet('可选灰度（新）：机台本地提供 GrayOptIn 参与开关；云端策略 requireOptIn=true 时仅 opt-in 机台参与百分比灰度，白名单可配置豁免，关键机台永不被强制升级。')
bullet('向后兼容：旧版 Launcher（不识别灰度字段）继续按 latest 逻辑工作，混布期间不产生故障；灰度收敛后（percent=100 且版本并入 latest）体系自然回归单一版本。')
bullet('全程可观测：机台通过现有可写 SFTP 账号周期上报心跳（机器标识、当前版本、灰度命中状态、opt-in 状态），运维在服务器端 heartbeat\\ 目录即可汇总查看灰度进度。')
bullet('可回滚：灰度期内发现问题，运维修改配置即可停止放量；已升级机台由 Launcher 在检测到目标版本低于本地版本时自动经 backup\\ 降级回滚，单机异常由崩溃恢复链路兜底。')
bullet('最小侵入：更新下载链路维持 SMB 只读不变；心跳上报复用主程序现有 SftpService 与 DPAPI 凭据体系；Launcher 传输层零改动。')

heading('1.3 范围与技术约束', 2)
tbl(['约束项', '内容', '对方案的影响'], [
    ['服务器形态', '更新服务器为纯 SMB 文件共享，无后端服务，不能做服务端动态下发', '灰度判定必须放在客户端，服务器只维护声明式策略文件'],
    ['latest 兼容', '旧版 Launcher 只认 version.json 的 latest 字段', 'gray 字段必须为可选扩展，旧客户端读不到时忽略并走原逻辑'],
    ['SMB 权限', '机台对更新共享只有读权限，且不计划开放写权限', '更新下载（只需读）维持现状；心跳写入不能走 SMB'],
    ['SFTP 账号', '已确认：每台机台均至少配置了一个可写 SFTP 账号（Role 非 readonly）', '心跳上报统一走 SFTP，复用现有 SftpService 与凭据'],
    ['SMB 多身份限制', 'Windows 同一会话对同一服务器仅允许一套身份（错误 1219）', '不引入 SMB 专用账号，规避 1219；更新检查直接使用 UNC 路径读取，不做 net use 挂载'],
    ['运行环境', 'WPF + .NET Framework 4.7.2，Launcher 与主程序同机同 Windows 用户运行', 'DPAPI（CurrentUser + 固定 entropy）加密的凭据两进程均可解密，可复用'],
    ['业务隔离', '心跳与更新检查为运维手段，绝不能影响文件搬运业务', '心跳失败静默重试，不弹窗、不阻塞；运行中后台更新检查不干扰 UI 与搬运；未登录传输账号时降级不上报'],
    ['检查时机（已定稿）', 'Launcher 启动时同步检查（6 秒超时兜底）+ 主程序运行中每 4 小时后台检查一次', '修正 1.1.2 风险1"从不检查更新"；放量后无需等机台重启即可被感知，下次启动生效'],
])

heading('1.4 关键决策记录', 2)
tbl(['决策点', '结论', '理由'], [
    ['更新下载通道', '维持 SMB 只读', '下载本身只需读权限，机台现状即满足；改造为 SFTP 下载需在 Launcher 引入 SSH.NET 并重写安装流程，收益低风险高'],
    ['心跳上报通道', 'SFTP（复用现有可写账号）', '每台机台均有可写 SFTP 账号（已确认）；SftpService.UploadFile/EnsureDirectoryExists 现成；无 SMB 1219 多身份问题'],
    ['灰度判定位置', '客户端自主判定（确定性分桶）', '服务器无后端服务，只能下发静态策略；哈希分桶保证同一机器判定结果稳定、放量只增不减'],
    ['机器标识', '首次生成 GUID 落盘（machine.id），此后复用', '主机名可能重名或被修改，不能作为稳定标识'],
    ['可选灰度（新）', '本地 config.xml 增加 GrayOptIn 开关（默认 false）+ 云端 gray.requireOptIn 策略位；白名单命中默认豁免 opt-in', '满足现场"自愿参与"诉求；默认不参与保证关键机台安全；白名单豁免保证内测机必中、不依赖现场操作'],
    ['检查时机（新）', '启动时检查 + 主程序运行中每 4 小时后台检查', '修正现行 Launcher"本地可用即跳过检查"导致从不更新的问题；6 秒超时兜底保证弱网可启动'],
    ['版本真相源（新）', '仅以 versions\\ 目录名为准，包内嵌 version.txt 交叉校验；禁止再使用程序集版本号做更新比较', '1.1.1 根因1：AssemblyInfo 写死 1.0.3.0 发版不同步导致无限提示'],
    ['安装执行器（新）', '废弃 update.bat，改为 Launcher --apply 安装模式（同步代码执行、逐步写 update.log）', '1.1.1 根因2：bat 编码地雷 + 失败无人值守无日志'],
    ['下载完整性（新）', '版本目录携带 manifest.json（文件清单 + SHA256），下载后逐文件校验通过才允许安装', '杜绝半成品安装；SMB 复制中断可检出'],
    ['失败兜底（新）', '同一版本安装失败 ≥2 次写入 update.failed，跳过该版本并经心跳上报，人工介入前不再尝试；同版本用户跳过后不重复弹窗', '1.1.2 风险3：状态机死循环；旧故障"一直提示"'],
    ['Launcher 自更新（新）', '版本包可携带 Launcher.new，主程序启动时检测并在 Launcher 未运行窗口期替换', '1.1.2 风险5：灰度判定逻辑需触达全部机台'],
    ['心跳连接策略', '复用主程序已登录的活跃 SFTP 连接，未登录时不上报', '避免额外握手；覆盖灰度最关键观测点——升级后主程序是否存活、版本是否正确'],
    ['密码存储', '复用 DpapiHelper（DPAPI/CurrentUser + 固定 entropy）', 'Launcher 与主程序同机同用户，可直接解密 config.xml 中的既有凭据，不新增凭据体系'],
])

# ===================================================================
# 第二章 初步思路与计划
# ===================================================================
heading('第二章 初步思路与计划', 1)

heading('2.1 总体架构（形态1+）', 2)
p('形态1+ 的核心分工：SMB 只负责"下发"（版本文件、清单与灰度策略，机台只读即可），SFTP 只负责"回收"（心跳上报，机台已有可写账号）。两条通道物理分离，互不干扰：')
code(r'''
+------------------ SMB 服务器（只读） ------------------+
|  version.json          <- latest + gray 灰度策略块      |
|  versions\1.0.6\        <- 灰度版本文件                  |
|  versions\1.0.6\manifest.json  <- 文件清单+SHA256(新增) |
+---------------------------+-----------------------------+
                            | 读策略 / 读版本文件（UNC 直读，不挂载）
+---------------------------+-----------------------------+
|  客户端 Launcher（传输层零改动）                          |
|  启动时检查（6s 超时，超时直接启动旧版）                   |
|  1. machineId  = MachineIdStore.GetOrCreate()           |
|  2. optIn      = config.xml GrayOptIn（本地可选灰度）     |
|  3. decision   = GrayPolicy.Decide(machineId, optIn,    |
|                                    gray)                |
|     优先级：黑名单 > 白名单(豁免opt-in) > requireOptIn    |
|             过滤 > 百分比哈希分桶                          |
|  4. 命中 -> 下载+manifest校验 -> Launcher --apply 安装    |
|     未命中 -> 走 latest（与现状一致）                      |
|  5. 写 gray.state（判定轨迹）+ update.log（安装日志）     |
|  6. 同版本失败>=2次 -> update.failed 熔断，不再尝试        |
+---------------------------+-----------------------------+
                            | 启动主程序, 携带 gray.state
+---------------------------+-----------------------------+
|  主程序 MoveImageForm                                    |
|  UpdateCheckService（新增）：运行中每 4 小时后台查一次     |
|  HeartbeatService（新增，基于 SftpService 实现）           |
|  定期上传 {machineId, sftpUser, hostname, version,       |
|            grayStatus, optIn, time}                     |
|  首启时检测版本包内 Launcher.new -> 替换 Launcher.exe     |
+---------------------------+-----------------------------+
                            | 上传 heartbeat\{machineId}.json
+---------------------------+-----------------------------+
|  SFTP 服务器                                             |
|  heartbeat\        <- 新建目录，各机台账号可写            |
|  （运维在此汇总查看灰度进度与异常熔断机台）                  |
+---------------------------------------------------------+
''')
p('说明：SFTP 服务器上的 heartbeat\\ 为一次性运维配置（对现有可写账号开放写权限即可）；业务目录结构与现有部署完全不变。', color=GRAY, size=9.5)

heading('2.2 核心设计要点', 2)

heading('2.2.1 version.json 灰度扩展（向后兼容）', 3)
code(r'''
{
  "latest": "1.0.5",
  "gray": {
    "version":   "1.0.6",
    "percent":   20,
    "whitelist": ["6f1c2b8e-...."],          // machineId 列表
    "blacklist": ["a1b2c3d4-...."],
    "requireOptIn": true,                    // 新增：百分比灰度仅对 opt-in 机台生效
    "whitelistOverridesOptIn": true,         // 新增：白名单命中豁免 opt-in（默认 true）
    "startTime": "2026-10-09T09:00:00"
  },
  "versions": {
    "1.0.5": { "date": "...", "note": "..." },
    "1.0.6": { "date": "...", "note": "..." }
  }
}
''')
bullet('gray 为可选字段：旧版 Launcher 解析不到时直接忽略，继续走 latest，天然向后兼容。')
bullet('放量操作 = 运维手工修改 percent 数值（20 → 50 → 100），改文件即生效，无需重启任何服务。')
bullet('灰度版本文件须先上传至 versions\\{gray.version}\\ 目录，包含 manifest.json，并同时在 versions 块中登记版本说明。')
bullet('全量收敛：percent 提至 100 且观察稳定后，将 latest 改为该版本并删除 gray 块，体系回归单一版本。')

heading('2.2.2 客户端灰度判定（GrayPolicy，确定性分桶 + opt-in 过滤）', 3)
p('判定优先级固定为：黑名单 > 白名单 > opt-in 过滤 > 百分比。百分比不能使用随机数（每次启动结果不同，无法收敛），必须使用稳定哈希分桶：')
code(r'''
// GrayPolicy.Decide 核心逻辑（示意）
if (gray.blacklist.Contains(machineId)) return GrayDecision.Stay;      // 强制排除
if (gray.whitelist.Contains(machineId)
    && gray.whitelistOverridesOptIn)  return GrayDecision.Upgrade;     // 白名单豁免 opt-in
if (gray.requireOptIn && !optIn)      return GrayDecision.Stay;        // 可选灰度：未报名不参与
int bucket = StableHash(machineId) % 100;   // FNV-1a 等稳定哈希，不引入新依赖
return bucket < gray.percent
    ? GrayDecision.Upgrade                  // 升级到 gray.version
    : GrayDecision.Stay;                    // 维持 latest 逻辑
''')
bullet('确定性：同一 machineId 永远落入同一桶，percent 由 20 调至 50 时原先命中的机器仍然命中，满足灰度"放量只增不减"语义。')
bullet('opt-in 单调性：机台从 opt-out 改为 opt-in 后，若其分桶在当前 percent 范围内，下一次检查即命中升级；反之 opt-in 改回 opt-out 后，已升级的灰度版本不自动降级（避免反复横跳），由运维决定是否摘除。')
bullet('哈希算法采用 FNV-1a 等内置可实现的稳定哈希，不引入 NuGet 依赖，保持 .NET Framework 4.7.2 零新依赖。')

heading('2.2.3 机器标识（MachineIdStore）', 3)
p('Launcher 首次启动生成 GUID，落盘至应用根目录 machine.id 文件，此后每次启动复用同一标识。封装为 MachineIdStore.GetOrCreate() 静态方法，供灰度判定与心跳上报共同使用（心跳文件名即 heartbeat\\{machineId}.json，一机一文件、覆盖写、不堆积）。白名单/黑名单中登记的即为 machineId。')

heading('2.2.4 可选灰度（本地 GrayOptIn 开关）', 3)
p('机台级"是否参与灰度"的自愿开关，落实为两层配置：')
tbl(['层', '配置', '语义'], [
    ['本地（机台）', 'config.xml 新增 <GrayOptIn>false</GrayOptIn>；主程序"账号管理"或设置区提供"参与灰度测试"复选框', '默认 false 不参与；现场/运维可手动开启。该值随心跳上报，运维可核对哪些机台已报名'],
    ['云端（策略）', 'version.json gray.requireOptIn（默认 true）', 'true 时百分比灰度仅覆盖 opt-in 机台；false 时百分比灰度覆盖全部机台（恢复强制放量语义，用于灰度后期扩量）'],
    ['豁免', 'gray.whitelistOverridesOptIn（默认 true）', '白名单命中的机台无视 opt-in 直接升级，保证内测机不依赖现场操作必中'],
])
p('判定矩阵（gray.version = 1.0.6，percent = 20，requireOptIn = true）：')
tbl(['机台情形', '判定结果', '说明'], [
    ['在 blacklist', 'Stay（latest）', '黑名单最高优先级'],
    ['在 whitelist（豁免开启）', 'Upgrade → 1.0.6', '内测机必中'],
    ['opt-in 且分桶 < 20', 'Upgrade → 1.0.6', '自愿报名且被百分比命中'],
    ['opt-in 且分桶 ≥ 20', 'Stay（latest）', '报名但未命中，等放量'],
    ['未 opt-in', 'Stay（latest）', 'requireOptIn=true 时不参与百分比灰度'],
])

heading('2.2.5 更新检查时机（启动时 + 运行中定时）', 3)
p('现行 Launcher "本地版本可用即跳过检查"导致现网从不检查更新（1.1.2 风险1），定稿为双通道检查：')
tbl(['通道', '触发', '行为', '兜底'], [
    ['启动时检查', 'Launcher 每次启动', 'UNC 直读 version.json（不做 net use 挂载）→ 灰度判定 → 有更新且未熔断则提示/安装', '6 秒超时或任何异常 → 直接启动本地现有版本，绝不阻塞开机自启链路'],
    ['运行中检查', '主程序 UpdateCheckService，每 4 小时一次', '后台线程读 version.json 并判定，结果写 gray.state（pending 标记）；命中灰度在下次启动由 Launcher 安装', '失败静默，下周期重试；不弹窗、不影响搬运；仅"有可用更新且未提示过"时在日志区提示一条'],
])

heading('2.2.6 更新链路可靠性设计（避开旧故障的六条硬措施）', 3)
tbl(['#', '旧坑（对应 1.1.1/1.1.2）', '新设计'], [
    ['1', '程序集版本号当真相，发版不同步', '版本真相源只用 versions\\ 目录名；包内嵌 version.txt 交叉校验；构建脚本自动注入版本号，杜绝人工遗忘'],
    ['2', 'update.bat 编码地雷、失败无人值守', '废弃 bat，改 Launcher --apply 安装模式：同步代码执行复制/校验/备份/回滚，每步写 update.log，结果写入 update.status（done/failed+原因）'],
    ['3', '下载不校验，半成品直接装', '版本目录携带 manifest.json（文件清单 + SHA256），下载到 versions\\.temp 后逐文件校验，全部通过才置 ready'],
    ['4', '失败死循环、一直提示', '同一版本安装失败 ≥2 次写 update.failed 熔断，跳过该版本并经心跳上报；记录 dismissedVersion，用户跳过的版本本次不再弹窗'],
    ['5', '升级"成功"无法确认', '安装完成三重确认：目录存在 + manifest 校验通过 + 主程序 --version 探测进程返回目标版本；任一失败自动从 backup\\ 回滚'],
    ['6', 'Launcher 自身不可更新', '版本包可携带 Launcher.new；主程序启动时检测并在 Launcher 未运行时替换为 Launcher.exe（替换前备份 Launcher.old）'],
])

heading('2.2.7 灰度轨迹（gray.state）', 3)
p('Launcher/主程序完成判定后写入本地 gray.state 纯文本文件，内容形如：')
code(r'''
decision=upgrade     # upgrade | stay | rollback | pending | failed
target=1.0.6         # 本次目标版本
source=gray          # gray | latest
optIn=true           # 本机可选灰度开关状态
time=2026-10-09T09:00:12
policyPercent=20
''')
p('主程序 HeartbeatService 读取该文件，与自身实际运行版本一并上报，形成"判定 → 升级 → 运行"的完整可观测链路：运维既能看到策略命中情况，也能确认升级后主程序是否真正活着；decision=failed 的机台即为熔断机台，需人工介入。')

heading('2.2.8 心跳上报（HeartbeatService，SFTP 实现）', 3)
code(r'''
{
  "machineId":  "6f1c2b8e-....",
  "sftpUser":   "vp_upload_01",
  "hostname":   "PC-A1",
  "version":    "1.0.6",
  "grayStatus": "applied",      // eligible | applied | stay | rollback | failed
  "optIn":      true,           // 新增：本机灰度参与开关
  "time":       "2026-10-09T09:05:00"
}
''')
bullet('实现完全复用 SftpService：EnsureDirectoryExists("heartbeat") + UploadFile，上传即三行调用；失败返回 TransferUploadResult.Fail 由上层静默处理。')
bullet('账号选择：从当前已登录的活跃 SFTP Profile 中取第一个 Role 非 readonly 的账号（心跳需要写权限）；可在 config.xml 增加 HeartbeatProfileName 显式指定优先账号。')
bullet('连接策略：复用主程序已登录的活跃连接（不额外握手）；未登录任何 SFTP 账号时本周期不上报，等待下个周期。')
bullet('节流设计：沿用现有传输心跳的 Stopwatch 双阈值模式（时间间隔 5 分钟 / 事件触发即拍），启动后首拍立即上报（覆盖"升级后是否活着"关键观测点）。')
bullet('失败语义：任何异常静默捕获、记入本地 Logs\\、下个周期重试；心跳永不弹窗、永不阻塞搬运业务。')

heading('2.2.9 回滚机制', 3)
tbl(['场景', '触发方式', '行为'], [
    ['灰度期内发现问题（批量止损）', '运维删除 version.json 的 gray 块，或将 gray.version 改回旧版本号', '未升级机台立即停止放量；已升级机台的 Launcher 检测到目标版本低于本地版本时，经 backup\\ 执行降级替换（同样过 manifest 校验与 --apply 日志）'],
    ['单机升级后启动异常（自救）', '安装后三重确认任一失败，或 VerifyCurrentVersion 校验失败', '自动从 backup\\ 回滚最近可用版本，写 update.log 与 gray.state decision=rollback，无需人工在场'],
    ['同版本反复安装失败（熔断）', '失败计数 ≥2', '写 update.failed 跳过该版本，心跳上报 failed，运维人工排查后删除该文件解除熔断'],
    ['机台退出灰度（opt-out）', '现场将 GrayOptIn 改为 false', '已升级的灰度版本不自动降级（避免横跳）；如需回退，由运维将其 machineId 加入 blacklist 并配合降级发布'],
])

heading('2.3 五阶段实施计划总览', 2)
tbl(['阶段', '内容', '主要产出', '风险'], [
    ['阶段0 更新链路可靠性修复', '版本真相源改造、manifest 校验、废弃 bat 改 --apply、失败熔断、检查时机恢复、Launcher 自更新', 'Launcher 1.0.7（无灰度能力，先行全量替换）、构建脚本自动注入版本号', '中（动更新主流程，但是后续一切的地基）'],
    ['阶段1 灰度判定核心库', 'GrayPolicy（含 opt-in 维度）+ MachineIdStore 纯逻辑实现与单元测试', 'Launcher/GrayPolicy.cs、Launcher/MachineIdStore.cs、测试工程', '低（纯逻辑，不动现有代码）'],
    ['阶段2 Launcher 集成', 'CheckForUpdate 接入灰度判定与 requireOptIn 过滤、写 gray.state、降级回滚分支、更新面板文案', 'Launcher 1.1.0（灰度能力）', '中（关键路径，依赖阶段0的可靠链路）'],
    ['阶段3 主程序心跳与运行中检查', 'HeartbeatService（SFTP）、UpdateCheckService（4h 后台检查）、AppConfig 扩展（GrayOptIn 等）、设置界面"参与灰度测试"开关', 'MoveImageForm/Services/HeartbeatService.cs、UpdateCheckService.cs、主程序新版', '中（动主程序，需验证不影响业务）'],
    ['阶段4 SOP 与闭环验证', '发布手册、测试环境全流程演练、混布验证', '灰度发布操作手册、验证报告、正式发布', '低（文档 + 测试为主）'],
])
p('实施顺序严格按 0 → 1 → 2 → 3 → 4：阶段0 是地基——不先修复更新链路，灰度只会放大旧故障；阶段1 纯逻辑可先行单测；阶段2 是灰度核心能力落地；阶段3 提供可观测性与运行中感知；阶段4 收尾固化运维流程。', bold=True)

# ===================================================================
# 第三章 实施内容与过程
# ===================================================================
heading('第三章 实施内容与过程', 1)

heading('3.1 阶段0：更新链路可靠性修复（先行）', 2)
p('目标：在引入任何灰度逻辑之前，把"检查 → 下载 → 安装 → 确认 → 回滚"链路修到可靠，消除 1.1.1 全部三类故障与 1.1.2 五项风险。本阶段不包含灰度能力，产出可作为一次普通版本全量发布（顺带验证新链路）。')
heading('3.1.1 任务分解', 3)
tbl(['任务', '改动位置', '说明', '验收标准'], [
    ['版本真相源统一', 'Launcher CheckForUpdate/GetLocalLatestVersion、构建脚本', '本地版本只取 versions\\ 目录名最大值且要求目录内含 manifest.json 校验通过；包内嵌 version.txt；MSBuild 后事件自动写入版本号', '构造 versions\\1.0.6 半成品目录（缺文件）时 Launcher 判为不可用并回退到 1.0.5 启动'],
    ['manifest 清单与校验', '发布打包脚本、Launcher 下载流程', '打包时生成 manifest.json（相对路径 + SHA256 + 大小）；下载到 versions\\.temp 后逐文件校验，任一不符删除 .temp 判失败', '人为篡改 .temp 中一个 dll 后，安装被拒绝并记 update.log'],
    ['废弃 bat，改 --apply 安装', 'Launcher', 'Launcher.exe --apply <version> 由启动方拉起：执行备份→复制→校验→探活→回滚（如需），全程写 update.log；update.status 终态为 done/failed', '安装中途杀进程，下次启动能从 backup\\ 自愈；update.log 记录完整步骤'],
    ['失败熔断与提示去重', 'Launcher', 'failed 计数写入 update.status 侧车文件；同版本 ≥2 次失败写 update.failed 跳过；dismissedVersion 记录用户跳过版本', '连续两次制造安装失败后出现 update.failed，Launcher 不再尝试该版本且正常启动旧版'],
    ['检查时机恢复', 'Launcher Window_Loaded', '移除"本地可用即跳过"分支，启动时必查（6 秒超时）；超时/异常直接启动本地版', '断网启动 Launcher：6 秒内进入主程序；联网启动：能发现新版本'],
    ['Launcher 自更新', '主程序启动逻辑、版本包规范', '版本包含 Launcher.new 时，主程序启动早期检测并替换（先备份 Launcher.old，替换失败还原）', '旧 Launcher 机台安装含 Launcher.new 的版本包后，下次启动 Launcher 已为新版本'],
])
heading('3.1.2 关键产出', 3)
code(r'''
Launcher/
  Program.cs / App.xaml.cs   // 支持 --apply <version> 命令行安装模式
  UpdateInstaller.cs         // 备份/复制/校验/探活/回滚 + update.log + 失败计数
  ManifestVerifier.cs        // manifest.json 解析与 SHA256 逐文件校验
tools/
  pack_release.py            // 打包：自动注入版本号、生成 manifest.json、可选携带 Launcher.new
''')

heading('3.2 阶段1：灰度判定核心库', 2)
p('目标：以纯逻辑类完成灰度判定全部规则（含 opt-in 维度），不触碰任何现有代码，配套单元测试先行验证正确性（TDD）。')
heading('3.2.1 任务分解', 3)
tbl(['任务', '说明', '验收标准'], [
    ['新建 GrayConfig 模型', 'version/percent/whitelist/blacklist/requireOptIn/whitelistOverridesOptIn/startTime 七字段，JSON 反序列化，全部可选', 'gray 块缺失或字段缺失时构造出的对象不触发灰度；requireOptIn 缺省视为 true、whitelistOverridesOptIn 缺省视为 true'],
    ['实现 GrayPolicy.Decide()', '黑名单 > 白名单(豁免) > requireOptIn 过滤 > 百分比四级判定，FNV-1a 稳定哈希分桶', '单测覆盖：优先级顺序、opt-in 开/关两种情形、白名单豁免、同 machineId 分桶稳定、percent=0/100 边界、空名单'],
    ['实现 MachineIdStore', 'GetOrCreate()：存在即读、缺失即生成 GUID 并落盘 machine.id', '单测覆盖：首次生成、二次读取一致、文件损坏时重建'],
    ['放量单调性验证', '模拟 1000 个 machineId，percent 20→50→100 逐级判定', '每一级命中集合严格包含上一级（只增不减）'],
    ['opt-in 语义验证', 'requireOptIn=true 时 opt-out 机台恒为 Stay；whitelistOverridesOptIn=true 时白名单无视 opt-in', '判定矩阵 2.2.4 五行全部有对应用例'],
])
heading('3.2.2 关键产出代码结构', 3)
code(r'''
Launcher/
  GrayPolicy.cs        // enum GrayDecision { Stay, Upgrade }
                       // static GrayDecision Decide(string machineId, bool optIn, GrayConfig gray)
  MachineIdStore.cs    // static string GetOrCreate(string appRoot)
''')

heading('3.3 阶段2：Launcher 集成', 2)
p('目标：把灰度判定接入 Launcher 更新主流程（基于阶段0修复后的可靠链路），这是灰度能力的核心落地阶段。')
heading('3.3.1 任务分解', 3)
tbl(['任务', '改动位置', '说明'], [
    ['解析 gray 字段', 'CheckForUpdate()', '读取 version.json 时同时解析 gray 块得到 GrayConfig；解析失败按无灰度处理'],
    ['读取本地 opt-in', 'Launcher 读 config.xml', '解析 <GrayOptIn>（默认 false）传入 GrayPolicy.Decide'],
    ['接入灰度判定', 'CheckForUpdate()', '目标版本 = Decide 为 Upgrade 时取 gray.version，否则取 latest；后续比较/下载/安装复用阶段0链路'],
    ['写 gray.state', '判定完成后', '按 2.2.7 格式写入本地（含 optIn 字段），供主程序心跳读取'],
    ['降级回滚分支', 'CheckForUpdate() 版本比较处', '目标版本 < 本地版本时，从服务器 versions\\ 复制旧版并经 --apply 替换（manifest 校验 + backup\\），写 gray.state decision=rollback'],
    ['更新面板文案', 'MainWindow UI', '区分"灰度更新（v1.0.6，灰度 20%，本机已报名参与）"与"正式更新"；熔断版本显示"该版本已暂停更新，请联系运维"'],
    ['旧字段兼容回归', '整体', 'gray 块不存在时行为与阶段0完全一致（走 latest），作为回归基准'],
])
heading('3.3.2 验收标准', 3)
bullet('测试服务器配置 gray{version=1.0.6, percent=20, requireOptIn=true, whitelist=[测试机machineId]}：白名单机器（无论是否 opt-in）升级 1.0.6；opt-in 且命中的机器升级；未 opt-in 的机器维持 1.0.5。')
bullet('requireOptIn 改为 false 后：未 opt-in 但分桶命中的机器也开始升级。')
bullet('删除 gray 块后重启 Launcher：已升级到 1.0.6 的机器自动降级回 1.0.5（经 backup\\），降级过程 update.status 状态机与 update.log 正确记录。')
bullet('旧版 Launcher（1.0.7，无灰度能力）指向含 gray 块的 version.json：行为与无 gray 时完全一致，无报错。')

heading('3.4 阶段3：主程序心跳上报（SFTP）与运行中检查', 2)
p('目标：为灰度过程提供可观测性，并让不放心的现场可以"看见"灰度进度；运行中定时检查让放量无需等待机台重启。')
heading('3.4.1 任务分解', 3)
tbl(['任务', '说明'], [
    ['新建 HeartbeatService.cs', '位于 MoveImageForm/Services/；读取 gray.state + 当前版本 + machineId + optIn，构造 JSON，经活跃 SFTP 连接上传 heartbeat\\{machineId}.json（覆盖写）'],
    ['新建 UpdateCheckService.cs', '运行中每 4 小时后台读 version.json 判定，命中写 gray.state（pending）；仅首次发现时在日志区提示"检测到新版本 vX.X.X，将在下次启动时更新"；失败静默'],
    ['账号选择逻辑', '优先 config.xml 的 HeartbeatProfileName（新增可选项）；否则取第一个已登录且 Role 非 readonly 的 SFTP Profile；无可用账号本周期跳过'],
    ['AppConfig 扩展', '新增 GrayOptIn（默认 false）、HeartbeatEnabled（默认 true）、HeartbeatIntervalMinutes（默认 5）、HeartbeatProfileName（可选）、UpdateCheckIntervalHours（默认 4）五个 XML 配置项'],
    ['设置界面开关', '账号管理 Tab 增加"参与灰度测试"复选框，绑定 GrayOptIn，改动即存 config.xml'],
    ['定时器与节流', 'DispatcherTimer 或后台 Task 周期触发；主程序启动后 30 秒内首拍心跳；沿用 Stopwatch 双阈值节流模式'],
    ['降级策略', 'SFTP 上传失败（含未登录、readonly 账号、网络异常）时：写本地 Logs\\heartbeat.log 一条记录，静默等待下周期；连续失败不累计告警、不影响 UI'],
    ['登出清理', '会话登出断开 SFTP 连接时，心跳随之暂停（无活跃连接即不上报），重新登录后自动恢复'],
])
heading('3.4.2 验收标准', 3)
bullet('灰度命中机器升级后：SFTP 服务器 heartbeat\\ 下出现对应 machineId 的 JSON，version=1.0.6、grayStatus=applied、optIn=true，且随周期刷新 time 字段。')
bullet('熔断机器（update.failed 存在）的心跳 grayStatus=failed，运维可据此拉清单人工排查。')
bullet('断开 SFTP 服务器网络：主程序搬运业务与 UI 正常，无弹窗无卡顿，本地 Logs\\ 留有失败记录，恢复网络后心跳自动续传。')
bullet('仅有 readonly 账号登录的会话：不上传、不报错、本地日志记录跳过原因。')
bullet('运维修改 percent 后：运行中的机器在下一个 4 小时检查点写入 gray.state pending，重启后完成升级（无需人工逐台通知重启时间）。')

heading('3.5 阶段4：运维 SOP 与闭环验证', 2)
heading('3.5.1 灰度发布标准操作流程（SOP）', 3)
code(r'''
0. 地基：阶段0版本（Launcher 1.0.7）已全量部署，更新链路可靠
1. 准备：pack_release.py 打包新版（如 1.0.6，自动注入版本号+生成 manifest），
   上传至服务器 versions\1.0.6\，并在 version.json 的 versions 块登记说明
   （此时尚不改 latest）
2. 内测：设置 gray = { version:1.0.6, percent:0, requireOptIn:true,
   whitelist:[测试机machineId] } —— 白名单单机验证（豁免 opt-in）
3. 小流量：观察 heartbeat 确认测试机 applied 且存活，
   将 percent 提至 20（仅 opt-in 机台参与；可同时清空 whitelist）
4. 放量：观察 24-48 小时（heartbeat 无 failed/rollback、版本分布符合预期），
   需要扩大覆盖时可将 requireOptIn 改为 false 放开全部机台，
   percent 提至 50，再观察，提至 100
5. 全量收敛：latest 改为 1.0.6，删除 gray 块，
   下个周期全部机台回归统一 latest 逻辑
6. 回滚（如需）：直接删除 gray 块或将 gray.version 改回旧版本，
   已升级机台由 Launcher 自动降级，观察 heartbeat 出现 rollback→stay
7. 熔断处理：heartbeat 中 grayStatus=failed 的机台人工排查，
   修复后删除该机 update.failed 文件解除熔断
''')
heading('3.5.2 测试环境演练清单', 3)
tbl(['场景', '步骤', '预期'], [
    ['更新链路可靠性', '断网启动 / 篡改下载文件 / 安装中途杀进程 / 连续两次安装失败', '6s 超时正常起旧版 / 校验拒绝 / backup 自愈 / update.failed 熔断不再重试'],
    ['白名单灰度', 'percent=0 + 白名单1台（豁免 opt-in）', '仅白名单机器升级，其余不动'],
    ['可选灰度', 'requireOptIn=true，两台分桶命中机台一开一关 opt-in', '仅 opt-in 机台升级；opt-out 机台维持 latest'],
    ['放开强制放量', 'requireOptIn 改 false', '未 opt-in 的命中机台下一次检查后开始升级'],
    ['百分比放量', 'percent=20，模拟多机', '约20%命中，命中集合稳定'],
    ['放量递增', '20→50→100', '命中集合逐级严格扩大，无机器被摘除'],
    ['黑名单', '将已命中机器加入 blacklist', '该机器维持旧版；若已升级，配合降级发布回退'],
    ['灰度回滚', '删除 gray 块', '已升级机器自动降级，heartbeat 显示 rollback'],
    ['旧版混布', '阶段0 Launcher + 含 gray 的 version.json', '行为与无 gray 完全一致'],
    ['运行中检查', '机台不重启，运维改 percent', '4 小时内 gray.state 出现 pending，重启后升级'],
    ['心跳链路', '灰度机器升级后', 'heartbeat\\ 出现该机 JSON 且周期刷新，optIn 字段正确'],
    ['心跳容错', '断开 SFTP 网络', '业务无影响，恢复后续传'],
])

heading('3.6 风险识别与应对', 2)
tbl(['风险', '影响', '应对'], [
    ['阶段0 改造引入更新主流程缺陷', '高：可能导致机台无法启动主程序', '阶段0 单独发布并先在测试环境全场景演练；backup\\ + 崩溃恢复兜底；保留旧 Launcher 安装包可人工回退'],
    ['灰度版本本身有缺陷', '中：波及灰度命中机器', '白名单单机先行；requireOptIn=true 期间仅自愿机台参与，天然限定爆炸半径；观察期不足不放量；删 gray 块即止血'],
    ['opt-in 机台数量过少', '低：百分比灰度覆盖不足，观察样本不够', '心跳上报 optIn 分布，运维可视报名情况决定何时将 requireOptIn 改为 false 放开'],
    ['现场误关 opt-in 后要求回退', '低：单台版本不一致', 'opt-out 不自动降级（避免横跳）；运维加 blacklist + 降级发布定向回退'],
    ['SFTP 账号现场差异（个别机器仅有 readonly）', '低：该机心跳缺失，灰度进度统计不全', 'HeartbeatService 跳过 readonly 并记本地日志；此类机器靠抽查版本号兜底；HeartbeatProfileName 支持显式指定'],
    ['heartbeat\\ 目录被写满或被篡改', '低：观测数据失真', '一机一文件覆盖写不堆积；服务器端对该目录设磁盘配额与定期清理；目录只写不删亦可'],
    ['machine.id 被误删或复制镜像部署', '低：分桶漂移、心跳文件重复', 'machine.id 缺失即重建（等效新机器，最多多观察一轮）；部署规范中注明勿整盘克隆 client 目录'],
    ['DPAPI 换机/换用户后无法解密', '低：心跳无可用账号', '现有机制：上层检测到空密码提示重新输入；心跳等待重新登录后自然恢复'],
    ['运行中检查与启动检查并发', '低：gray.state 写入竞争', 'gray.state 写入采用临时文件+原子替换；主程序只写 pending 不触发安装，安装动作恒由 Launcher 执行'],
])

heading('3.7 里程碑与交付节点', 2)
tbl(['里程碑', '内容', '交付物'], [
    ['M0 更新链路可靠（新）', '阶段0 完成，测试环境全场景演练通过，Launcher 1.0.7 全量替换', 'Launcher 1.0.7、UpdateInstaller/ManifestVerifier、pack_release.py、演练记录'],
    ['M1 判定库完成', '阶段1 完成，单测全绿', 'GrayPolicy.cs、MachineIdStore.cs、测试工程与用例'],
    ['M2 灰度能力上线（测试环境）', '阶段2 完成，测试环境通过 3.3.2 验收', 'Launcher 1.1.0、灰度回滚分支、面板文案'],
    ['M3 可观测性上线', '阶段3 完成，通过 3.4.2 验收', 'HeartbeatService、UpdateCheckService、AppConfig 扩展、主程序新版'],
    ['M4 定稿发布', '阶段4 演练全通过，SOP 评审', '灰度发布操作手册、验证报告、Deploy-Server 更新源更新'],
])

# ===================================================================
# 第四章 最终期望结果
# ===================================================================
heading('第四章 最终期望结果', 1)

heading('4.1 功能验收标准', 2)
tbl(['验收项', '标准'], [
    ['更新链路可靠（新）', '下载经 manifest 校验、安装有 update.log 与终态、失败自动回滚、同版本失败熔断、同版本不重复提示；1.1.1 三类故障在测试中不可复现'],
    ['检查时机（新）', 'Launcher 每次启动必查（6s 超时兜底可启动）；主程序运行中每 4 小时后台检查，放量后无需人工逐台重启'],
    ['灰度判定', '黑名单 > 白名单(豁免) > opt-in 过滤 > 百分比判定正确；同一机器分桶结果永久稳定；percent 递增时命中集合只增不减'],
    ['可选灰度（新）', 'requireOptIn=true 时仅 opt-in 机台参与百分比灰度；白名单命中豁免；opt-in 状态随心跳上报可查'],
    ['放量控制', '运维仅通过修改服务器 version.json 即可完成 0% → 20% → 50% → 100% 的逐级放量，改文件即生效'],
    ['向后兼容', '旧版 Launcher 在含 gray 配置的环境下行为与从前完全一致；灰度收敛后体系回归 latest 单一逻辑'],
    ['心跳可观测', '服务器端 heartbeat\\ 可实时汇总：每台机器的 machineId、版本、灰度状态、opt-in、最后心跳时间；异常机器（无心跳/rollback/failed）可被识别'],
    ['回滚能力', '删除 gray 块即可批量止损；已升级机器自动降级；单机启动异常自动从 backup\\ 自救'],
    ['业务零影响', '心跳与后台检查任何失败不影响文件搬运业务与 UI；灰度升级过程对操作员可见（面板文案）但不改变操作习惯'],
])

heading('4.2 交付物清单', 2)
tbl(['类别', '交付物'], [
    ['源代码', 'Launcher --apply 安装模式、ManifestVerifier、UpdateInstaller、GrayPolicy.cs、MachineIdStore.cs、Launcher 灰度集成、MoveImageForm/Services/HeartbeatService.cs、UpdateCheckService.cs、AppConfig 扩展、设置界面 opt-in 开关'],
    ['工具', 'pack_release.py（打包自动注入版本号 + 生成 manifest.json + 可选携带 Launcher.new）'],
    ['测试', '更新链路故障注入用例（篡改/中断/杀进程/连续失败）、灰度判定单元测试（分桶稳定性、优先级、opt-in 矩阵、边界）、测试环境演练记录'],
    ['部署', 'Launcher 1.0.7 / 1.1.0 安装包、主程序新版安装包、Deploy-Server 更新源（version.json 含 gray 规范）、SFTP 服务器 heartbeat\\ 目录配置说明'],
    ['文档', '灰度发布操作手册（SOP）、version.json gray 字段规范、本规划方案书'],
])

heading('4.3 典型使用场景（发布演练剧本）', 2)
p('以 1.0.6 版本发布为例（阶段0 地基已全量部署），完整生命周期如下：')
code(r'''
Day 1  10:00  pack_release.py 打包 1.0.6（自动注入版本号+manifest），
              上传 versions\1.0.6\，versions 块登记说明（不改 latest）
       10:30  gray: percent=0, requireOptIn=true, whitelist=[测试机]
              → 测试机升级（豁免 opt-in）
       11:00  查 heartbeat：测试机 applied、心跳正常、搬运业务正常
Day 1  14:00  percent=20 → 已 opt-in 的机台中约20%分桶命中升级
Day 2  09:00  查 heartbeat：全部命中机器 applied、无 rollback/failed、心跳连续
       （若异常：删除 gray 块 → 自动降级 → 修复后重新开始；
        failed 机台人工排查后删 update.failed 解除熔断）
Day 3  09:00  requireOptIn=false，放开全部机台参与百分比灰度
Day 3  14:00  percent=50 → 观察一轮（运行中机器 4h 内感知，重启后升级）
Day 4  09:00  percent=100 → 全量灰度
Day 5  09:00  latest=1.0.6，删除 gray 块 → 收敛完成
       09:30  heartbeat 全部机器 version=1.0.6、grayStatus 归于 stay
''')
p('至此，项目达成"链路可靠、小步放量、自愿参与、全程可观测、随时可回滚"的灰度更新能力：发布风险从"一次全量、出错全灭"收敛为"逐级验证、可控止损"，且 1.1.1 的旧更新故障在机制上被彻底消除。', bold=True)

# ===================================================================
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   'VP运维工具_灰度更新项目规划方案书.docx')
doc.save(out)
print('saved:', out)
